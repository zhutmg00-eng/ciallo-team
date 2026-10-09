import os
import sys
import re
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
base = r'D:\谷雨光合计划校园挑战赛\文献资料'
files = [
    '00-说明.docx',
    '01-赛题简报.docx',
    '02-谷雨品牌资料.docx',
    '03-学术文献.docx',
    '04-行业报告与法规.docx',
    '05-同类比赛获奖案例.docx',
    '光合计划-科技AI赛道-文献资料汇编.docx'
]

url_pattern = re.compile(r'https?://[^\s\u4e00-\u9fa5\)\"\'\<\>]+')

total_items = {}

for fname in files:
    path = os.path.join(base, fname)
    if not os.path.exists(path):
        continue
    doc = Document(path)
    urls = []
    # paragraphs
    for p in doc.paragraphs:
        matches = url_pattern.findall(p.text)
        for m in matches:
            urls.append(m.rstrip(';,，。；'))
    # tables
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                matches = url_pattern.findall(cell.text)
                for m in matches:
                    urls.append(m.rstrip(';,，。；'))
    unique_urls = list(dict.fromkeys(urls))
    total_items[fname] = unique_urls
    print(f"{fname}: {len(urls)} total URLs, {len(unique_urls)} unique URLs")
