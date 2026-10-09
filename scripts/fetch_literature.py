import os
import sys
import re
import time
import json
import urllib.request
import urllib.error
from concurrent.futures import ThreadPoolExecutor, as_completed
from docx import Document
from lxml import html

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r'D:\谷雨光合计划校园挑战赛'
DOCS_DIR = os.path.join(BASE_DIR, '文献资料')
OUTPUT_DIR = os.path.join(BASE_DIR, '文献精读')
os.makedirs(OUTPUT_DIR, exist_ok=True)

USER_AGENT = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
url_pattern = re.compile(r'https?://[^\s\u4e00-\u9fa5\)\"\'\<\>]+')

def clean_url(u):
    return u.rstrip(';,，。；）)]}>"\'')

def extract_all_targets():
    targets = []
    
    # 1. Academic papers (03-学术文献.docx)
    p3 = os.path.join(DOCS_DIR, '03-学术文献.docx')
    if os.path.exists(p3):
        doc3 = Document(p3)
        for t_idx, t in enumerate(doc3.tables):
            theme_name = f"学术文献_主题{t_idx + 1}"
            for r_idx, r in enumerate(t.rows[1:]):
                c = [cell.text.strip().replace('\n', ' ') for cell in r.cells]
                if len(c) >= 8:
                    num, title, author, year, journal, link_cell, concl, usage = c[0], c[1], c[2], c[3], c[4], c[5], c[6], c[7]
                    urls = url_pattern.findall(link_cell)
                    for u in urls:
                        targets.append({
                            'category': '03_学术文献',
                            'folder': theme_name,
                            'id': num,
                            'title': title,
                            'author': author,
                            'year': year,
                            'journal': journal,
                            'url': clean_url(u),
                            'context_summary': concl,
                            'context_usage': usage,
                            'priority': 1
                        })

    # 2. Industry reports & benchmarks (04-行业报告与法规.docx)
    p4 = os.path.join(DOCS_DIR, '04-行业报告与法规.docx')
    if os.path.exists(p4):
        doc4 = Document(p4)
        current_sec = "行业与法规"
        current_id = ""
        current_title = ""
        current_text = []
        for p in doc4.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if text.startswith(('A.', 'B.', 'C.', 'D.', '0.')) and len(text) < 50:
                current_sec = text[:30]
            elif re.match(r'^[A-D]\d+\.', text):
                current_id = text.split('.')[0]
                current_title = text
                current_text = [text]
            else:
                current_text.append(text)
                urls = url_pattern.findall(text)
                for u in urls:
                    prio = 1 if current_id.startswith('B') or current_id in ['A1', 'A9', 'A10', 'A12', 'A13', 'C1', 'C2', 'C6'] else 2
                    targets.append({
                        'category': '04_行业与法规',
                        'folder': f"行业法规_{current_sec.split('.')[0]}",
                        'id': current_id or 'General',
                        'title': current_title or text[:40],
                        'author': '',
                        'year': '',
                        'journal': '',
                        'url': clean_url(u),
                        'context_summary': " ".join(current_text[-3:])[:300],
                        'context_usage': '',
                        'priority': prio
                    })

    # 3. Winning cases (05-同类比赛获奖案例.docx)
    p5 = os.path.join(DOCS_DIR, '05-同类比赛获奖案例.docx')
    if os.path.exists(p5):
        doc5 = Document(p5)
        current_sec = "获奖案例"
        current_case = ""
        current_text = []
        for p in doc5.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if text.startswith(('A.', 'B.', 'C.', 'D.')) and len(text) < 50:
                current_sec = text[:30]
            elif '案例' in text and ('｜' in text or '|' in text):
                current_case = text
                current_text = [text]
            else:
                current_text.append(text)
                urls = url_pattern.findall(text)
                for u in urls:
                    targets.append({
                        'category': '05_获奖案例',
                        'folder': f"获奖案例_{current_sec.split('.')[0]}",
                        'id': current_case[:20] if current_case else 'Case',
                        'title': current_case or text[:40],
                        'author': '',
                        'year': '',
                        'journal': '',
                        'url': clean_url(u),
                        'context_summary': " ".join(current_text[-3:])[:300],
                        'context_usage': '',
                        'priority': 1
                    })

    # 4. Brand data (02-谷雨品牌资料.docx)
    p2 = os.path.join(DOCS_DIR, '02-谷雨品牌资料.docx')
    if os.path.exists(p2):
        doc2 = Document(p2)
        current_sec = "品牌资料"
        current_item = ""
        current_text = []
        for p in doc2.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if text.startswith(('一、', '二、', '三、', '四、', '五、', '六、', '七、', '八、', '九、', '十、', '十一、', '十二、')):
                current_sec = text[:30]
            elif re.match(r'^\d+\.', text):
                current_item = text
                current_text = [text]
            else:
                current_text.append(text)
                urls = url_pattern.findall(text)
                for u in urls:
                    targets.append({
                        'category': '02_品牌资料',
                        'folder': f"品牌资料_{current_sec[:6]}",
                        'id': current_item[:20] if current_item else 'Brand',
                        'title': current_item or text[:40],
                        'author': '',
                        'year': '',
                        'journal': '',
                        'url': clean_url(u),
                        'context_summary': " ".join(current_text[-3:])[:300],
                        'context_usage': '',
                        'priority': 1 if any(k in current_sec for k in ['概况', '营收', '光甘草定', '痛点', '数字化', '机会点']) else 2
                    })

    # Deduplicate by URL
    seen = set()
    unique_targets = []
    for t in targets:
        u = t['url']
        if u not in seen and u.startswith('http'):
            seen.add(u)
            unique_targets.append(t)
            
    return unique_targets

def fetch_target(item):
    url = item['url']
    req = urllib.request.Request(
        url,
        headers={
            'User-Agent': USER_AGENT,
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8'
        }
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            content_type = resp.headers.get('Content-Type', '')
            if 'pdf' in content_type.lower() or url.endswith('.pdf'):
                return item, False, "[PDF Document - binary format]", "PDF"
            raw_bytes = resp.read()
            encoding = 'utf-8'
            if 'charset=' in content_type:
                encoding = content_type.split('charset=')[-1].split(';')[0].strip()
            elif b'charset="gbk"' in raw_bytes[:1024].lower() or b'charset=gbk' in raw_bytes[:1024].lower() or b'charset=gb2312' in raw_bytes[:1024].lower():
                encoding = 'gbk'
            try:
                html_text = raw_bytes.decode(encoding, errors='replace')
            except Exception:
                html_text = raw_bytes.decode('utf-8', errors='replace')
            
            tree = html.fromstring(html_text)
            for el in tree.xpath('//script | //style | //nav | //footer | //header | //aside'):
                el.getparent().remove(el)
            
            title = ""
            t_elems = tree.xpath('//title')
            if t_elems and t_elems[0].text:
                title = t_elems[0].text.strip()
                
            article_elems = tree.xpath('//article | //main | //div[contains(@class, "article") or contains(@class, "content") or contains(@class, "post") or contains(@class, "main")]')
            if article_elems:
                main_elem = article_elems[0]
            else:
                main_elem = tree.xpath('//body')[0] if tree.xpath('//body') else tree
                
            raw_text = main_elem.text_content()
            lines = [l.strip() for l in raw_text.split('\n') if l.strip()]
            cleaned_text = '\n\n'.join(lines)
            return item, True, cleaned_text[:12000], title
    except Exception as e:
        return item, False, str(e), ""

def sanitize_filename(name):
    cleaned = re.sub(r'[\\/*?:"<>|\r\n\t]', '_', str(name)).strip(' .')
    return cleaned[:50] if cleaned else "unnamed"

def main():
    targets = extract_all_targets()
    print(f"Total targets: {len(targets)}")
    
    # Sort by priority
    targets.sort(key=lambda x: x['priority'])
    
    # Let's crawl in parallel
    print("Starting parallel fetch for literature...")
    results_summary = []
    success_count = 0
    fail_count = 0
    
    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(fetch_target, t): t for t in targets}
        for future in as_completed(futures):
            item, success, content, page_title = future.result()
            subfolder = os.path.join(OUTPUT_DIR, sanitize_filename(item['folder']))
            os.makedirs(subfolder, exist_ok=True)
            
            base_name = f"{sanitize_filename(str(item['id']))}_{sanitize_filename(item['title'])}"
            md_path = os.path.join(subfolder, f"{base_name}.md")
            
            if success and len(content) > 100:
                success_count += 1
                with open(md_path, 'w', encoding='utf-8') as out_f:
                    out_f.write(f"# {item['title']}\n\n")
                    out_f.write(f"- **ID / 序号**: {item['id']}\n")
                    out_f.write(f"- **类别**: {item['category']} / {item['folder']}\n")
                    out_f.write(f"- **URL**: {item['url']}\n")
                    out_f.write(f"- **网页标题**: {page_title}\n")
                    if item['author']:
                        out_f.write(f"- **作者/机构**: {item['author']}\n")
                    if item['journal'] or item['year']:
                        out_f.write(f"- **期刊/年份**: {item['journal']} ({item['year']})\n")
                    out_f.write(f"\n## 资料包原生注释\n- **核心结论/摘要**: {item['context_summary']}\n")
                    if item['context_usage']:
                        out_f.write(f"- **赛题应用价值**: {item['context_usage']}\n")
                    out_f.write(f"\n## 抓取全文/正文内容\n\n{content}\n")
                
                results_summary.append({
                    'id': item['id'],
                    'title': item['title'],
                    'url': item['url'],
                    'status': 'success',
                    'chars': len(content),
                    'path': md_path
                })
            else:
                fail_count += 1
                results_summary.append({
                    'id': item['id'],
                    'title': item['title'],
                    'url': item['url'],
                    'status': 'failed',
                    'error': content[:100]
                })

    print(f"\nFetch completed! Success: {success_count}, Failed/PDF/Blocked: {fail_count}")
    with open(os.path.join(BASE_DIR, 'fetch_summary.json'), 'w', encoding='utf-8') as sf:
        json.dump(results_summary, sf, ensure_ascii=False, indent=2)

if __name__ == '__main__':
    main()
